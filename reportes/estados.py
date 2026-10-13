"""Estados financieros y formatos del Ciclo 5 — Información Gerencial (Fase 10).

PDF (reportlab):
* Balance General, Estado de Resultados, Estado de Flujo de Efectivo (método
  directo) y Estado de Evolución del Patrimonio Neto con la estructura del
  MAN-006 §18, y Notas 1 a 16 (§19); paquete completo en un solo PDF.
* FIN-INF-01 (Lista de Verificación de Cierre Mensual), FIN-INF-03
  (Consolidado de Costeo Mensual), FIN-INF-04 (Estado de Resultados Mensual e
  Indicadores), INF-CV-04 (Informe Previo de Fiscalización) y la liquidación
  individual del socio (CI-FIN-02 Art. 8).

Excel (openpyxl): paquete de estados con una hoja por estado y FIN-INF-04.

Las funciones reciben los diccionarios de ``servicios.cierre`` y devuelven
bytes. No importan ``streamlit``.
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import PageBreak, Spacer, Table, TableStyle

import config
from reportes import pdf_contable as pc
from reportes.formatos_tesoreria import (
    E_NORMAL, E_PEQ, E_SUB, E_TITULO, VERSION_FORMATO, Contexto, _bs, _contexto, _firmas, _p, _tabla,
)
from utils.formato import fecha_pantalla

INFO = "Ciclo 5 — Información Gerencial (Consolidación)"
PIE = "Ciclo 5 — Información Gerencial"
MESES = ("ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE",
         "NOVIEMBRE", "DICIEMBRE")
SEMAFORO_TEXTO = {"BUENO": "Bueno", "NORMAL": "Normal", "MALO": "Malo", None: "Sin ventas"}
E_NEGRITA = ParagraphStyle("pl_negrita_estados", parent=E_NORMAL, fontName="Helvetica-Bold")
E_CENTRO = ParagraphStyle("pl_centro_estados", parent=E_NORMAL, alignment=1)
E_PEQ_CENTRO = ParagraphStyle("pl_peq_centro_estados", parent=E_PEQ, alignment=1)

__all__ = ["Contexto", "balance_pdf", "resultados_pdf", "flujo_pdf", "patrimonio_pdf", "notas_pdf", "paquete_pdf",
           "fin_inf_01", "fin_inf_03", "fin_inf_04", "inf_cv_04", "liquidacion_socio", "paquete_excel",
           "fin_inf_04_excel", "numero"]

_R12 = ["12", "Cierre mensual con comprobantes no registrados o mal clasificados, distorsionando el ICL, el RINC y el "
        "ROA reportados a Gerencia.", "Alto", "Preventivo", "Clasificación"]
_R32 = ["32", "Manipulación de los indicadores de gestión para mostrar cumplimiento artificial de metas frente al "
        "Consejo de Administración.", "Moderado", "Detección", "Verificación y Evaluación"]
_R04 = ["4", "Pago de excedentes que deteriore el Saldo Mínimo de Caja del ciclo siguiente.", "Alto", "Preventivo",
        "Autorización / Protección Física"]
_R30 = ["30", "Distribución de excedentes sin constituir previamente las reservas legales obligatorias (Legal 5%, "
        "Educación 1%).", "Moderado", "Preventivo", "Autorización"]


def numero(valor: Any) -> str:
    """Monto para estados: ``1.234,56`` y los negativos entre paréntesis."""
    if valor is None or valor == "":
        return ""
    d = Decimal(str(valor))
    texto = f"{abs(d):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"({texto})" if d < 0 else texto


def _fecha_larga(iso: str) -> str:
    dia, mes, anio = iso[8:10], int(iso[5:7]), iso[:4]
    return f"{int(dia)} DE {MESES[mes - 1]} DE {anio}"


def _cabecera_estado(titulo: str, subtitulo: str, con_titulo: bool = False) -> list[Any]:
    """Cabecera del estado dentro del cuerpo. En un PDF de un solo estado el
    título va en el membrete de cada página; en el paquete se repite aquí."""
    cuerpo: list[Any] = []
    if con_titulo:
        cuerpo += [_p(titulo, E_TITULO), _p(subtitulo, E_CENTRO)]
    return cuerpo + [_p("(Expresado en Bolivianos)", E_PEQ_CENTRO), Spacer(1, 6)]


_ESTILO_FILA = {
    "titulo": ("Helvetica-Bold", pc.VERDE_CLARO, 0),
    "subtitulo": ("Helvetica-Bold", None, 0.3),
    "total": ("Helvetica-Bold", pc.VERDE_TOTAL, 0),
    "resultado": ("Helvetica-Bold", pc.VERDE_TOTAL, 0),
    "texto": ("Helvetica-Oblique", None, 0),
    "linea": ("Helvetica", None, 0.6),
}


def _tabla_estado(filas: Sequence[Mapping[str, Any]], con_notas: bool = False) -> Table:
    """Estado con el estilo institucional: encabezado verde, bordes, montos a la
    derecha (negativos entre paréntesis) y totales destacados."""
    datos = [["Concepto", "Nota", "Bs"] if con_notas else ["Concepto", "Bs"]]
    estilos: list[Any] = [("BACKGROUND", (0, 0), (-1, 0), pc.VERDE), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                          ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.3),
                          ("ALIGN", (-1, 0), (-1, -1), "RIGHT"), ("ALIGN", (1, 0), (1, -1), "CENTER"),
                          ("BOX", (0, 0), (-1, -1), 0.8, pc.VERDE), ("INNERGRID", (0, 0), (-1, -1), 0.25, pc.GRIS_LINEA),
                          ("TOPPADDING", (0, 0), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]
    if not con_notas:
        estilos.remove(("ALIGN", (1, 0), (1, -1), "CENTER"))
    for i, f in enumerate(filas, start=1):
        fuente, fondo, sangria = _ESTILO_FILA.get(f["tipo"], _ESTILO_FILA["linea"])
        concepto = ("    " * int(sangria * 2)) + f["concepto"]
        fila = [concepto, f.get("nota", ""), numero(f["monto"])] if con_notas else [concepto, numero(f["monto"])]
        datos.append(fila)
        estilos.append(("FONTNAME", (0, i), (-1, i), fuente))
        if fondo is not None:
            estilos.append(("BACKGROUND", (0, i), (-1, i), fondo))
        if f["tipo"] in ("total", "resultado"):
            estilos.append(("LINEABOVE", (0, i), (-1, i), 0.8, pc.VERDE))
    anchos = [12.6 * cm, 1.8 * cm, 4.0 * cm] if con_notas else [14.4 * cm, 4.0 * cm]
    tabla = Table(datos, colWidths=anchos, repeatRows=1)
    tabla.setStyle(TableStyle(estilos))
    return tabla


def _control(texto: str, cuadra: bool) -> Any:
    return pc.verificacion(texto, cuadra)


def _espacio_notas() -> list[Any]:
    """Referencia a las notas y espacio para observaciones (MAN-006 §19)."""
    caja = Table([[_p("Notas y observaciones:", E_NEGRITA)], [""], [""], [""]], colWidths=[18.4 * cm],
                 rowHeights=[14, 14, 14, 14])
    caja.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, pc.GRIS_LINEA),
                              ("LINEBELOW", (0, 1), (-1, 2), 0.3, pc.GRIS_LINEA)]))
    return [Spacer(1, 6), _p("Las Notas 1 a 16 adjuntas son parte integrante e inseparable de este estado y deben "
                             "leerse conjuntamente con él (MAN-006 §19).", E_PEQ), Spacer(1, 3), caja]


def _cuerpo_balance(bg: Mapping[str, Any], con_titulo: bool = False) -> list[Any]:
    return _cabecera_estado("BALANCE GENERAL", f"AL {_fecha_larga(bg['fecha_corte'])}", con_titulo) + [
        _tabla_estado(bg["filas"], con_notas=True), Spacer(1, 4),
        _control(f"Activo {_bs(bg['total_activo'])} = Pasivo + Patrimonio "
                 f"{_bs(Decimal(str(bg['total_pasivo'])) + Decimal(str(bg['total_patrimonio'])))}", bg["cuadra"])]


def _subtitulo_resultados(er: Mapping[str, Any]) -> str:
    return (f"Por la gestión {er['ejercicio']} que inicia el 01 de julio (del {fecha_pantalla(er['desde'])} al "
            f"{fecha_pantalla(er['hasta'])})")


def _cuerpo_resultados(er: Mapping[str, Any], con_titulo: bool = False) -> list[Any]:
    return _cabecera_estado("ESTADO DE RESULTADOS", _subtitulo_resultados(er), con_titulo) + [
        _tabla_estado(er["filas"]), Spacer(1, 4),
        _control(f"Ingresos totales {_bs(er['total_ingresos'])}; excedente neto {_bs(er['excedente_neto'])}",
                 er["cuadra"])]


def _subtitulo_flujo(efe: Mapping[str, Any]) -> str:
    return f"Por el periodo comprendido al {_fecha_larga(efe['hasta'])} (Método Directo)"


def _cuerpo_flujo(efe: Mapping[str, Any], con_titulo: bool = False) -> list[Any]:
    return _cabecera_estado("ESTADO DE FLUJO DE EFECTIVO", _subtitulo_flujo(efe), con_titulo) + [
        _tabla_estado(efe["filas"]), Spacer(1, 4),
        _control(f"Saldo final {_bs(efe['saldo_final'])} frente a disponibilidades según libros "
                 f"{_bs(efe['disponibilidades_segun_libros'])}", efe["cuadra"])]


def _subtitulo_patrimonio(ee: Mapping[str, Any]) -> str:
    return f"Por la gestión {ee['ejercicio']} (del {fecha_pantalla(ee['desde'])} al {fecha_pantalla(ee['hasta'])})"


def _cuerpo_patrimonio(ee: Mapping[str, Any], con_titulo: bool = False) -> list[Any]:
    columnas = ee["columnas"]
    datos = [["Concepto"] + [n for _, n in columnas] + ["Total"]]
    for f in ee["filas"]:
        datos.append([f["concepto"]] + [numero(f[c]) for c, _ in columnas] + [numero(f["total"])])
    ancho_concepto = 5.4
    ancho = (18.4 - ancho_concepto) / (len(columnas) + 1)
    tabla = Table([[_p(c, pc.E_ENCAB) for c in datos[0]]] + datos[1:],
                  colWidths=[ancho_concepto * cm] + [ancho * cm] * (len(columnas) + 1), repeatRows=1)
    estilos = [("BACKGROUND", (0, 0), (-1, 0), pc.VERDE), ("FONTSIZE", (0, 1), (-1, -1), 8),
               ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("BOX", (0, 0), (-1, -1), 0.8, pc.VERDE),
               ("INNERGRID", (0, 0), (-1, -1), 0.3, pc.GRIS_LINEA), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for i, f in enumerate(ee["filas"], start=1):
        if f["tipo"] == "total":
            estilos += [("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"), ("BACKGROUND", (0, i), (-1, i), pc.VERDE_TOTAL)]
    tabla.setStyle(TableStyle(estilos))
    return _cabecera_estado("ESTADO DE EVOLUCIÓN DEL PATRIMONIO NETO", _subtitulo_patrimonio(ee), con_titulo) + [
        tabla, Spacer(1, 4), _control(f"Total {_bs(ee['total_final'])} igual al patrimonio del Balance General",
                                      ee["cuadra"])]


def _cuerpo_notas(notas: Sequence[Mapping[str, Any]], fecha_corte: str, con_titulo: bool = False) -> list[Any]:
    cuerpo: list[Any] = []
    if con_titulo:
        cuerpo += [_p(f"NOTAS A LOS ESTADOS FINANCIEROS AL {_fecha_larga(fecha_corte)}", E_TITULO), Spacer(1, 6)]
    cuerpo.append(_p("Las notas son parte integrante e inseparable del Balance General, del Estado de Resultados, del "
                     "Estado de Evolución del Patrimonio Neto y del Estado de Flujo de Efectivo, y deben leerse "
                     "conjuntamente con ellos.", E_NORMAL))
    for nota in notas:
        cuerpo.append(_p(f"NOTA {nota['numero']} — {nota['titulo']}", E_SUB))
        for b in nota["bloques"]:
            if b["tipo"] == "parrafo":
                cuerpo.append(_p(b["texto"], E_NORMAL))
                cuerpo.append(Spacer(1, 3))
            elif b["tipo"] == "subtitulo":
                cuerpo.append(_p(b["texto"], E_NEGRITA))
            else:
                n = len(b["encabezados"])
                if n == 2:
                    anchos = [13.0, 5.4]
                elif n == 3:
                    anchos = [10.0, 3.4, 5.0]
                else:
                    anchos = [18.4 * 0.34] + [18.4 * 0.66 / (n - 1)] * (n - 1)
                filas = list(b["filas"])
                totales = (len(filas) - 1,) if b.get("totales") else ()
                cuerpo += [pc.tabla(b["encabezados"], filas, anchos, montos=tuple(b["alinear"]), totales=totales),
                           Spacer(1, 4)]
    return cuerpo


FIRMAS_ESTADOS = ("Representante Legal — Gerente General", "Contador General — Unidad de Finanzas",
                  "Presidente del Consejo de Administración")


def _firmas_estados() -> list[Any]:
    return [pc.firmas(FIRMAS_ESTADOS), Spacer(1, 2),
            pc.firmas(("Vº Bº Consejo de Vigilancia (INF-CV-08)",))]


def _documento(titulo: str, subtitulo: str, codigo: str, cuerpo: list[Any], membrete: pc.Membrete | None,
               emision: pc.Emision | None) -> bytes:
    return pc.construir(titulo, subtitulo, codigo, cuerpo, membrete, emision)


def balance_pdf(bg: Mapping[str, Any], membrete: pc.Membrete | None = None,
                emision: pc.Emision | None = None) -> bytes:
    """Balance General (MAN-006 §18.1)."""
    return _documento("Balance General", f"Al {_fecha_larga(bg['fecha_corte']).lower()}", "MAN-006 §18.1",
                      _cuerpo_balance(bg) + _espacio_notas() + _firmas_estados(), membrete, emision)


def resultados_pdf(er: Mapping[str, Any], membrete: pc.Membrete | None = None,
                   emision: pc.Emision | None = None) -> bytes:
    """Estado de Resultados (MAN-006 §18.2)."""
    return _documento("Estado de Resultados", _subtitulo_resultados(er), "MAN-006 §18.2",
                      _cuerpo_resultados(er) + _espacio_notas() + _firmas_estados(), membrete, emision)


def flujo_pdf(efe: Mapping[str, Any], membrete: pc.Membrete | None = None,
              emision: pc.Emision | None = None) -> bytes:
    """Estado de Flujo de Efectivo, método directo (MAN-006 §18.3)."""
    return _documento("Estado de Flujo de Efectivo", _subtitulo_flujo(efe), "MAN-006 §18.3",
                      _cuerpo_flujo(efe) + _espacio_notas() + _firmas_estados(), membrete, emision)


def patrimonio_pdf(ee: Mapping[str, Any], membrete: pc.Membrete | None = None,
                   emision: pc.Emision | None = None) -> bytes:
    """Estado de Evolución del Patrimonio Neto (MAN-006 §18.4)."""
    return _documento("Estado de Evolución del Patrimonio Neto", _subtitulo_patrimonio(ee), "MAN-006 §18.4",
                      _cuerpo_patrimonio(ee) + _espacio_notas() + _firmas_estados(), membrete, emision)


def notas_pdf(notas: Sequence[Mapping[str, Any]], fecha_corte: str, membrete: pc.Membrete | None = None,
              emision: pc.Emision | None = None) -> bytes:
    """Notas 1 a 16 a los estados financieros (MAN-006 §19)."""
    return _documento("Notas a los Estados Financieros", f"Al {_fecha_larga(fecha_corte).lower()}", "MAN-006 §19",
                      _cuerpo_notas(notas, fecha_corte), membrete, emision)


def paquete_pdf(paquete: Mapping[str, Any], huella: str | None = None, membrete: pc.Membrete | None = None,
                emision: pc.Emision | None = None) -> bytes:
    """Los cuatro estados, indicadores y las Notas 1-16 en un solo PDF."""
    cuerpo = _cuerpo_balance(paquete["balance"], True) + [PageBreak()] + _cuerpo_resultados(paquete["resultados"], True)
    ind = paquete["indicadores"]
    filas = [[k, v["formula"], _pct_txt(v["valor"]), SEMAFORO_TEXTO.get(v["semaforo"], "—"),
              paquete.get("umbrales", {}).get(k, "")] for k, v in ind.items()]
    cuerpo += [Spacer(1, 8), _p("Indicadores de gestión del ejercicio (PR-FIN-10)", E_SUB),
               pc.tabla(["Indicador", "Fórmula", "Resultado", "Semáforo", "Umbral"], filas, [1.8, 5.6, 2.2, 2.2, 6.6],
                        montos=(2,))]
    cuerpo += [PageBreak()] + _cuerpo_flujo(paquete["flujo"], True) + [PageBreak()]
    cuerpo += _cuerpo_patrimonio(paquete["patrimonio"], True) + _firmas_estados()
    cuerpo += [PageBreak()] + _cuerpo_notas(paquete["notas"], paquete["fecha_corte"], True)
    if huella:
        cuerpo += [Spacer(1, 8), _p(f"Huella SHA-256 del paquete emitido: {huella}", E_PEQ)]
    return _documento("Estados Financieros", f"Ejercicio {paquete['ejercicio']} — al "
                      f"{_fecha_larga(paquete['fecha_corte']).lower()}", f"EEFF {paquete['ejercicio']}", cuerpo,
                      membrete, emision)


def _pct_txt(valor: Any) -> str:
    if valor is None:
        return "—"
    return f"{Decimal(str(valor)) * 100:.2f} %".replace(".", ",")


def _riesgos(nota: str, filas: Sequence[Sequence[str]]) -> list[Any]:
    return [Spacer(1, 8), _p("Riesgos Mitigados — Referencia al Mapa de Riesgos (MAN-FIN-001)", E_SUB),
            _p(nota, E_PEQ), Spacer(1, 3),
            _tabla(["N°", "Riesgo (síntesis)", "Nivel", "Tipo de Control", "Objetivo COSO"], filas,
                   [1, 8.6, 2.2, 2.6, 3])]


def _ficha_formato(codigo: str, area: str, ciclo: str) -> list[Any]:
    """Código, versión, área y ciclo del formato (el título va en el membrete)."""
    return [pc.ficha([("Código", codigo), ("Versión", VERSION_FORMATO), ("Área / Proceso", area),
                      ("Ciclo transaccional", ciclo)]), Spacer(1, 4)]


def _cuerpo_formato(codigo: str, titulo: str, ctx: Contexto, descripcion: str, area: str = "Unidad de Finanzas y "
                    "Comercialización", ciclo: str = INFO) -> list[Any]:
    return _ficha_formato(codigo, area, ciclo) + _contexto(ctx) + [
        _p("Cuerpo de Datos", E_SUB), _p(descripcion, E_NORMAL), Spacer(1, 4)]


def _formato(codigo: str, titulo: str, cuerpo: list[Any], membrete: pc.Membrete | None,
             emision: pc.Emision | None) -> bytes:
    """Formato FIN-xx / INF-CV-xx con el membrete institucional unificado."""
    return pc.construir(titulo, f"Formato {codigo} — versión {VERSION_FORMATO}", codigo, cuerpo, membrete, emision)


# ==========================================================================
# FIN-INF-01, FIN-INF-03, FIN-INF-04
# ==========================================================================
def fin_inf_01(lista: Mapping[str, Any], ctx: Contexto, verificado_por: str = "",
               membrete: pc.Membrete | None = None, emision: pc.Emision | None = None) -> bytes:
    """Lista de Verificación de Cierre Mensual (pendientes en cero)."""
    filas = [[f["verificacion"], "Sí" if f["cumple"] else "No", str(f["pendientes"])] for f in lista["filas"]]
    cuerpo = _cuerpo_formato("FIN-INF-01", "Lista de Verificación de Cierre Mensual", ctx,
                             f"Verifica que todos los comprobantes contables de {lista['periodo']} (tesorería, compras, "
                             "planillas, costeo, ventas) estén registrados y correctamente clasificados según el Plan "
                             "de Cuentas, antes de iniciar la consolidación.")
    cuerpo.append(_tabla(["Verificación", "Cumple (Sí/No)", "N° de Comprobantes Pendientes"], filas, [10.4, 3, 4], (2,)))
    detalle = [d for f in lista["filas"] for d in f["detalle"]]
    if detalle:
        cuerpo += [Spacer(1, 4), _p("Detalle de pendientes:", E_NORMAL)] + [_p(f"• {d}", E_PEQ) for d in detalle]
    cuerpo.append(_p(f"Total de comprobantes pendientes: {lista['pendientes']}.", E_NORMAL))
    cuerpo += _firmas(["Verificado por (Unidad de Finanzas)"], [verificado_por or ctx.elaborado_por])
    cuerpo += _riesgos("Primer filtro de integridad del cierre mensual. Ningún cierre mensual debe consolidarse "
                       "(FIN-INF-03/04) con comprobantes pendientes sin resolver.", [_R12, _R32])
    return _formato("FIN-INF-01", "Lista de Verificación de Cierre Mensual", cuerpo, membrete, emision)


def fin_inf_03(consolidado: Mapping[str, Any], ctx: Contexto, verificado_por: str = "",
               membrete: pc.Membrete | None = None, emision: pc.Emision | None = None) -> bytes:
    """Consolidado de Costeo Mensual."""
    filas = [[c["ciclo"], _bs(c["costo_total"]), f"{Decimal(str(c['kg_producidos'])):,.3f}".replace(",", "X")
              .replace(".", ",").replace("X", "."), _bs(c["costo_unitario"]) if c["costo_unitario"] else "—", "", ""]
             for c in consolidado["ciclos"]]
    filas.append(["Total del mes (mayor 6.1.02)", _bs(consolidado["costo_ciclos"]), "", "",
                  f"{Decimal(str(consolidado['kg_vendidos'])):,.3f}".replace(",", "X").replace(".", ",")
                  .replace("X", "."), _bs(consolidado["costo_ventas"])])
    cuerpo = _cuerpo_formato("FIN-INF-03", "Consolidado de Costeo Mensual", ctx,
                             f"Consolida el costo real de producción de cada ciclo cerrado en {consolidado['periodo']} "
                             "(Informe de Costo de Ciclo, FIN-TRA-06) y determina el Costo de Ventas del período.")
    cuerpo.append(_tabla(["Lote / Ciclo Cerrado en el Mes", "Costo Total del Ciclo (Bs)", "Kg de Producto Terminado",
                          "Costo Unitario Bs/kg", "Kg Vendidos en el Mes", "Costo de Ventas del Mes (Bs)"], filas,
                         [3.4, 3, 2.6, 2.4, 2.6, 3.4], (1, 2, 3, 4, 5)))
    cuerpo.append(_p(f"Costo de ventas: carne {_bs(consolidado['costo_ventas_carne'])}; alimento "
                     f"{_bs(consolidado['costo_ventas_alimento'])}.", E_PEQ))
    cuerpo += _firmas(["Elaborado por (Unidad de Finanzas)", "Verificado por (Gerencia General)"],
                      [ctx.elaborado_por, verificado_por])
    cuerpo += _riesgos("Base del cálculo del Margen Bruto y del ICL del mes.", [_R12, _R32])
    return _formato("FIN-INF-03", "Consolidado de Costeo Mensual", cuerpo, membrete, emision)


def fin_inf_04(informe: Mapping[str, Any], ctx: Contexto, membrete: pc.Membrete | None = None,
               emision: pc.Emision | None = None) -> bytes:
    """Estado de Resultados Mensual e Indicadores (ICL, RINC, ROA)."""
    mes, acu = informe["mes"], informe["acumulado"]
    filas = [[f["concepto"], _bs(f["monto"]), _bs(a["monto"])] for f, a in zip(mes["filas"], acu["filas"])]
    cuerpo = _cuerpo_formato("FIN-INF-04", "Estado de Resultados Mensual e Indicadores", ctx,
                             f"Estado de Resultados en formato escalonado (PC-08) de {informe['periodo']} y cálculo de "
                             "los tres indicadores de gestión clave, previo a su verificación por el Consejo de "
                             "Vigilancia (CI-FIN-18).")
    cuerpo.append(_tabla(["Concepto", "Mes (Bs)", "Acumulado del ejercicio (Bs)"], filas, [9.4, 4, 4], (1, 2)))
    cuerpo.append(Spacer(1, 6))
    ind = []
    for clave in ("ICL", "RINC", "ROA"):
        m, a = informe["indicadores_mes"][clave], informe["indicadores_acumulado"][clave]
        ind.append([clave, m["formula"], f"{_pct_txt(m['valor'])} ({SEMAFORO_TEXTO.get(m['semaforo'])})",
                    f"{_pct_txt(a['valor'])} ({SEMAFORO_TEXTO.get(a['semaforo'])})", informe["umbrales"][clave]])
    cuerpo.append(_tabla(["Indicador", "Fórmula", "Resultado del Mes", "Acumulado", "Umbral"], ind,
                         [1.5, 4.6, 3, 3, 5.3]))
    cuerpo.append(_p(f"Activos totales al cierre del mes (base del ROA): {_bs(informe['activo_total'])}.", E_PEQ))
    cuerpo += _firmas(["Elaborado por (Unidad de Finanzas)",
                       "Remitido, antes de su comunicación oficial, a (Consejo de Vigilancia — CI-FIN-18)"],
                      [ctx.elaborado_por, ""])
    cuerpo += _riesgos("Este formato se remite al Consejo de Vigilancia antes de comunicarse a la Gerencia como "
                       "reporte oficial (CI-FIN-18, Paso 1).", [_R12, _R32])
    return _formato("FIN-INF-04", "Estado de Resultados Mensual e Indicadores", cuerpo, membrete, emision)


# ==========================================================================
# INF-CV-04 Y LIQUIDACIÓN INDIVIDUAL
# ==========================================================================
def inf_cv_04(distribucion: Mapping[str, Any], ctx: Contexto, membrete: pc.Membrete | None = None,
              emision: pc.Emision | None = None) -> bytes:
    """Informe Previo de Fiscalización del Consejo de Vigilancia."""
    informe = distribucion["informe_cv_04"]
    filas = [[v["verificacion"], "Sí" if v["cumple"] else "No", v["observacion"]] for v in informe["verificaciones"]]
    cuerpo = _cuerpo_formato(
        "INF-CV-04", "Informe Previo de Fiscalización", ctx,
        f"{informe['codigo']}. Emitido por el Consejo de Vigilancia antes de la Asamblea General, verificando el "
        "cumplimiento de la Ley N° 356, de los límites de reserva obligatoria y de la política de Saldo Mínimo de Caja "
        f"(FICHA-FIN-02 / CI-FIN-02, Paso 3). Excedente neto del ejercicio {distribucion['ejercicio']}: "
        f"{_bs(distribucion['excedente'])}; excedente distribuible propuesto: {_bs(distribucion['distribuible'])}.",
        area="Consejo de Vigilancia", ciclo="Ciclo 1 — Tesorería (Financiamiento)")
    cuerpo.append(_tabla(["Verificación", "Cumple (Sí/No)", "Observaciones"], filas, [7, 2.6, 7.8]))
    favorable = informe["dictamen"] == "FAVORABLE"
    cuerpo.append(_p(f"Dictamen del Consejo de Vigilancia: Favorable [{'X' if favorable else ' '}]   Con observaciones "
                     f"[{' ' if favorable else 'X'}]", E_NORMAL))
    if informe.get("observaciones"):
        cuerpo.append(_p(f"Observaciones: {informe['observaciones']}", E_NORMAL))
    cuerpo += _firmas(["Elaborado por (Consejo de Vigilancia)", "Recibido por (Asamblea General)"],
                      [informe["elaborado_por"], ""])
    cuerpo += _riesgos("Control preventivo previo, condición indispensable para la aprobación de la distribución de "
                       "excedentes por la Asamblea (CI-FIN-02 Art. 13).", [_R04, _R30])
    return _formato("INF-CV-04", "Informe Previo de Fiscalización", cuerpo, membrete, emision)


def liquidacion_socio(distribucion: Mapping[str, Any], liquidacion: Mapping[str, Any], ctx: Contexto,
                      membrete: pc.Membrete | None = None, emision: pc.Emision | None = None) -> bytes:
    """Liquidación individual firmada por la Gerencia General (CI-FIN-02 Art. 8)."""
    cuerpo = _ficha_formato("LIQ-EXC", "Gerencia General", "Ciclo 1 — Tesorería (Financiamiento)") + _contexto(ctx)
    cuerpo += [
        _p("Datos del socio", E_SUB),
        _tabla(["Socio", "N.º", "C.I.", "Certificados"],
               [[liquidacion["socio"], str(liquidacion["socio_numero"]),
                 f"{liquidacion['ci']} {liquidacion['ci_expedido']}", str(liquidacion["certificados"])]],
               [8.4, 1.6, 4, 3.4]),
        Spacer(1, 6), _p("Liquidación", E_SUB),
        _tabla(["Concepto", "Bs"], [
            [f"Excedente neto del ejercicio {distribucion['ejercicio']}", _bs(distribucion["excedente"])],
            ["(–) Reserva Legal", _bs(distribucion["reserva_legal"])],
            ["(–) Fondo de Educación Cooperativa", _bs(distribucion["educacion"])],
            ["(–) Reserva de Reinversión y Expansión", _bs(distribucion["reinversion"])],
            ["(–) Capitalización (Aportes Adicionales)", _bs(distribucion["capitalizacion"])],
            ["(=) Excedente distribuible", _bs(distribucion["distribuible"])],
            [f"Criterio de reparto: {distribucion['criterio'].lower()} (base {liquidacion['base']})", ""],
            ["Monto bruto asignado al socio", _bs(liquidacion["bruto"])],
            ["(–) Deducciones aplicables (RC-IVA cuando corresponda)", _bs(liquidacion["deducciones"])],
            ["(=) MONTO NETO A COBRAR", _bs(liquidacion["neto"])],
        ], [12.4, 5], (1,)),
        Spacer(1, 4),
        _p(f"Acta de la Asamblea: {distribucion.get('acta_ref') or '—'}. Estado: {liquidacion['estado'].lower()}"
           + (f"; firmada por {liquidacion['firmada_por']}" if liquidacion.get("firmada_por") else "")
           + (f"; pagada el {fecha_pantalla(liquidacion['pagada_en'])}" if liquidacion.get("pagada_en") else "") + ".",
           E_PEQ),
        _p("El pago se ejecuta por transferencia bancaria o cheque nominativo, solo si el saldo bancario posterior se "
           "mantiene igual o superior al SME vigente (CI-FIN-02 Arts. 9 y 10).", E_PEQ),
    ]
    cuerpo += _firmas(["Gerencia General", "Socio (recibí conforme)"], [liquidacion.get("firmada_por") or "",
                                                                        liquidacion["socio"]])
    return _formato("LIQ-EXC", "Liquidación Individual de Excedentes", cuerpo, membrete, emision)


# ==========================================================================
# EXCEL
# ==========================================================================
_FMT = '#,##0.00;(#,##0.00);"-"'
_NEGRITA = Font(bold=True)
_RELLENO = PatternFill("solid", fgColor="E3EEE7")
_LINEA = Border(top=Side(style="thin", color="5D6A62"))


def _encabezado_hoja(ws: Any, titulo: str, subtitulo: str, columnas: int) -> int:
    for fila, texto in enumerate((config.NOMBRE_ENTIDAD.upper(), titulo, subtitulo, "(Expresado en Bolivianos)"),
                                 start=1):
        ws.cell(row=fila, column=1, value=texto).font = Font(bold=fila <= 2, size=12 if fila <= 2 else 10)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=columnas)
        ws.cell(row=fila, column=1).alignment = Alignment(horizontal="center")
    return 6


def _hoja_estado(wb: Workbook, nombre: str, titulo: str, subtitulo: str, filas: Sequence[Mapping[str, Any]],
                 con_notas: bool = False) -> None:
    ws = wb.create_sheet(nombre)
    columnas = 3 if con_notas else 2
    fila = _encabezado_hoja(ws, titulo, subtitulo, columnas)
    encabezado = ["Concepto", "Notas", "Bs"] if con_notas else ["Concepto", "Bs"]
    for c, texto in enumerate(encabezado, start=1):
        celda = ws.cell(row=fila, column=c, value=texto)
        celda.font, celda.fill = _NEGRITA, _RELLENO
    for f in filas:
        fila += 1
        sangria = 2 if f["tipo"] == "linea" else (1 if f["tipo"] == "subtitulo" else 0)
        ws.cell(row=fila, column=1, value=f["concepto"]).alignment = Alignment(indent=sangria)
        if con_notas:
            ws.cell(row=fila, column=2, value=f.get("nota") or None)
        if f["monto"] is not None:
            celda = ws.cell(row=fila, column=columnas, value=float(Decimal(str(f["monto"]))))
            celda.number_format = _FMT
            if f["tipo"] in ("total", "resultado"):
                celda.border = _LINEA
        if f["tipo"] in ("titulo", "subtitulo", "total", "resultado"):
            for c in range(1, columnas + 1):
                ws.cell(row=fila, column=c).font = _NEGRITA
        if f["tipo"] in ("titulo", "resultado"):
            for c in range(1, columnas + 1):
                ws.cell(row=fila, column=c).fill = _RELLENO
    ws.column_dimensions["A"].width = 62
    for c in range(2, columnas + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18


def paquete_excel(paquete: Mapping[str, Any]) -> bytes:
    """Libro Excel con Balance, Resultados, Flujo, EEPN, Notas e Indicadores."""
    wb = Workbook()
    wb.remove(wb.active)
    bg, er, efe, ee = paquete["balance"], paquete["resultados"], paquete["flujo"], paquete["patrimonio"]
    _hoja_estado(wb, "Balance General", "BALANCE GENERAL", f"AL {_fecha_larga(bg['fecha_corte'])}", bg["filas"], True)
    _hoja_estado(wb, "Estado de Resultados", "ESTADO DE RESULTADOS",
                 f"Por la gestión {er['ejercicio']} que inicia el 01 de julio", er["filas"])
    _hoja_estado(wb, "Flujo de Efectivo", "ESTADO DE FLUJO DE EFECTIVO",
                 f"Al {_fecha_larga(efe['hasta'])} (Método Directo)", efe["filas"])
    ws = wb.create_sheet("Evolución Patrimonio")
    columnas = ee["columnas"]
    fila = _encabezado_hoja(ws, "ESTADO DE EVOLUCIÓN DEL PATRIMONIO NETO", f"Gestión {ee['ejercicio']}",
                            len(columnas) + 2)
    for c, texto in enumerate(["Concepto"] + [n for _, n in columnas] + ["Total"], start=1):
        celda = ws.cell(row=fila, column=c, value=texto)
        celda.font, celda.fill = _NEGRITA, _RELLENO
    for f in ee["filas"]:
        fila += 1
        ws.cell(row=fila, column=1, value=f["concepto"])
        for c, (clave, _) in enumerate(columnas, start=2):
            ws.cell(row=fila, column=c, value=float(Decimal(str(f[clave])))).number_format = _FMT
        ws.cell(row=fila, column=len(columnas) + 2, value=float(Decimal(str(f["total"])))).number_format = _FMT
        if f["tipo"] == "total":
            for c in range(1, len(columnas) + 3):
                ws.cell(row=fila, column=c).font = _NEGRITA
    ws.column_dimensions["A"].width = 34
    for c in range(2, len(columnas) + 3):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws = wb.create_sheet("Notas")
    fila = _encabezado_hoja(ws, "NOTAS A LOS ESTADOS FINANCIEROS", f"AL {_fecha_larga(paquete['fecha_corte'])}", 7)
    for nota in paquete["notas"]:
        fila += 1
        ws.cell(row=fila, column=1, value=f"NOTA {nota['numero']} — {nota['titulo']}").font = _NEGRITA
        for b in nota["bloques"]:
            if b["tipo"] in ("parrafo", "subtitulo"):
                fila += 1
                celda = ws.cell(row=fila, column=1, value=b["texto"])
                celda.alignment = Alignment(wrap_text=True, vertical="top")
                ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
                ws.row_dimensions[fila].height = max(15, 15 * (len(b["texto"]) // 110 + 1))
                if b["tipo"] == "subtitulo":
                    celda.font = _NEGRITA
            else:
                fila += 1
                for c, texto in enumerate(b["encabezados"], start=1):
                    celda = ws.cell(row=fila, column=c, value=texto)
                    celda.font, celda.fill = _NEGRITA, _RELLENO
                for r in b["filas"]:
                    fila += 1
                    for c, texto in enumerate(r, start=1):
                        ws.cell(row=fila, column=c, value=texto)
        fila += 1
    ws.column_dimensions["A"].width = 60
    for c in range(2, 8):
        ws.column_dimensions[get_column_letter(c)].width = 20
    ws = wb.create_sheet("Indicadores")
    fila = _encabezado_hoja(ws, "INDICADORES DE GESTIÓN", f"Gestión {paquete['ejercicio']}", 5)
    for c, texto in enumerate(["Indicador", "Fórmula", "Resultado", "Semáforo", "Umbral"], start=1):
        celda = ws.cell(row=fila, column=c, value=texto)
        celda.font, celda.fill = _NEGRITA, _RELLENO
    for clave, v in paquete["indicadores"].items():
        fila += 1
        ws.cell(row=fila, column=1, value=clave)
        ws.cell(row=fila, column=2, value=v["formula"])
        celda = ws.cell(row=fila, column=3, value=None if v["valor"] is None else float(Decimal(str(v["valor"]))))
        celda.number_format = "0.00%"
        ws.cell(row=fila, column=4, value=SEMAFORO_TEXTO.get(v["semaforo"]))
        ws.cell(row=fila, column=5, value=paquete.get("umbrales", {}).get(clave, ""))
    for c, ancho in zip("ABCDE", (12, 44, 14, 12, 52)):
        ws.column_dimensions[c].width = ancho
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def fin_inf_04_excel(informe: Mapping[str, Any]) -> bytes:
    """FIN-INF-04 en Excel (mes y acumulado con indicadores)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "FIN-INF-04"
    fila = _encabezado_hoja(ws, "ESTADO DE RESULTADOS MENSUAL E INDICADORES (FIN-INF-04)", informe["periodo"], 3)
    for c, texto in enumerate(["Concepto", "Mes (Bs)", "Acumulado (Bs)"], start=1):
        celda = ws.cell(row=fila, column=c, value=texto)
        celda.font, celda.fill = _NEGRITA, _RELLENO
    for m, a in zip(informe["mes"]["filas"], informe["acumulado"]["filas"]):
        fila += 1
        ws.cell(row=fila, column=1, value=m["concepto"])
        for c, f in ((2, m), (3, a)):
            ws.cell(row=fila, column=c, value=float(Decimal(str(f["monto"])))).number_format = _FMT
        if m["tipo"] in ("total", "resultado"):
            for c in range(1, 4):
                ws.cell(row=fila, column=c).font = _NEGRITA
    fila += 2
    for c, texto in enumerate(["Indicador", "Mes", "Acumulado", "Umbral"], start=1):
        celda = ws.cell(row=fila, column=c, value=texto)
        celda.font, celda.fill = _NEGRITA, _RELLENO
    for clave in ("ICL", "RINC", "ROA"):
        fila += 1
        ws.cell(row=fila, column=1, value=clave)
        for c, origen in ((2, "indicadores_mes"), (3, "indicadores_acumulado")):
            v = informe[origen][clave]["valor"]
            ws.cell(row=fila, column=c, value=None if v is None else float(Decimal(str(v)))).number_format = "0.00%"
        ws.cell(row=fila, column=4, value=informe["umbrales"][clave])
    for c, ancho in zip("ABCD", (48, 18, 18, 50)):
        ws.column_dimensions[c].width = ancho
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()

