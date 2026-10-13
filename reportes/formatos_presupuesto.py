"""Formatos de la Fase 12: presupuesto, Consejo de Vigilancia e indicadores.

PDF: FIN-TES-09 (hoja de supuestos), FIN-TES-10 (presupuesto por ciclo),
FIN-TES-11 (modificación presupuestaria), FIN-TES-12 (ejecución y
variaciones), FIN-TES-13 (cuadro de mando), INF-CV-02, INF-CV-07 y FIN-INF-05.
Excel: Presupuesto Maestro 2026-2030 y FIN-TES-12.
"""
from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.platypus import Spacer

from reportes.formatos_tesoreria import E_PEQ, E_SUB, Contexto, _bs, _construir, _contexto, _encabezado, _firmas, _p, _tabla
from utils.formato import fecha_pantalla

CICLO_PRESU = "Ciclo 1 — Tesorería / Presupuesto"
CICLO_INF = "Ciclo 5 — Información Gerencial (Consolidación)"
NOMBRE_SEMAFORO = {"VERDE": "Verde", "AMARILLO": "Amarillo", "ROJO": "Rojo", None: "—"}


def pct(valor: Any, decimales: int = 1) -> str:
    """Fracción como porcentaje con coma decimal (0,6241 → 62,4 %)."""
    if valor is None or valor == "":
        return "—"
    q = Decimal(1).scaleb(-decimales)
    return f"{(Decimal(str(valor)) * 100).quantize(q, rounding=ROUND_HALF_UP)}".replace(".", ",") + " %"


def num(valor: Any, decimales: int = 2) -> str:
    if valor is None or valor == "":
        return "—"
    texto = f"{Decimal(str(valor)):,.{decimales}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def fin_tes_09(ejercicio: int, filas: Sequence[Mapping[str, Any]], ctx: Contexto) -> bytes:
    cuerpo = _encabezado("FIN-TES-09", f"Hoja de Supuestos Presupuestarios — Ejercicio {ejercicio}", "Unidad de Finanzas",
                         CICLO_PRESU) + _contexto(ctx)
    cuerpo += [_tabla(["Parámetro", "Unidad", "Valor base 2026", "Valor del ejercicio", "Fuente / respaldo"],
                      [[f["parametro"], f["unidad"], num(f["base"], 4) if f["base"] is not None else "—",
                        num(f["valor"], 4) if f["valor"] is not None else "—", f["fuente"]] for f in filas],
                      [4.4, 1.8, 2.4, 2.4, 6.4], (2, 3)),
               Spacer(1, 4), _p("Los valores se leen de la tabla de parámetros vigente por ejercicio (regla 2).", E_PEQ)]
    cuerpo += _firmas(["Elaborado por (Unidad de Finanzas)", "Parámetros técnicos (Unidad Técnica)",
                       "Revisado por (Gerencia General)"])
    return _construir("FIN-TES-09", cuerpo, CICLO_PRESU)


def fin_tes_10(datos: Mapping[str, Any], ctx: Contexto) -> bytes:
    cuerpo = _encabezado("FIN-TES-10", f"Presupuesto por Ciclo — Ciclo {datos['ciclo']} del ejercicio "
                         f"{datos['ejercicio']}", "Unidad de Finanzas", CICLO_PRESU) + _contexto(ctx)
    cuerpo += [_p(f"Ventana del ciclo: {fecha_pantalla(datos['desde'])} al {fecha_pantalla(datos['hasta'])}", E_SUB)]
    grupos: dict[str, list[list[Any]]] = {}
    for f in datos["filas"]:
        grupos.setdefault(f["grupo"], []).append(
            [f["nombre"], num(f["cantidad"], 3) if f["cantidad"] is not None else "—",
             num(f["precio"], 4) if f["precio"] is not None else "—", _bs(f["monto"]), _bs(f["acumulado"]),
             "Modificado" if f["ajuste"] else ""])
    filas = []
    for grupo, lineas in grupos.items():
        filas.append([_p(grupo, E_SUB), "", "", "", "", ""])
        filas += lineas
    filas.append([_p("CAJA", E_SUB), "", "", "", "", ""])
    filas.append(["Saldo mínimo de efectivo (SME)", "", "", _bs(datos["sme"]), "", ""])
    filas.append(["Saldo final proyectado (flujo anual prorrateado)", "", "", _bs(datos["saldo_final_proyectado"]), "", ""])
    cuerpo += [_tabla(["Elemento", "Cantidad", "Precio unitario", "Presupuesto del ciclo", "Acumulado del año",
                       "Observaciones"], filas, [5.2, 2, 2.2, 2.9, 2.9, 2.2], (1, 2, 3, 4))]
    cuerpo += _firmas(["Elaborado por (Unidad de Finanzas)", "Revisado por (Gerencia General)",
                       "Aprobado por (Consejo de Administración)"])
    return _construir("FIN-TES-10", cuerpo, CICLO_PRESU)


def fin_tes_11(m: Mapping[str, Any], ctx: Contexto) -> bytes:
    cuerpo = _encabezado("FIN-TES-11", f"Solicitud de Modificación Presupuestaria N.º {m['id']}", "Gerencia General",
                         CICLO_PRESU) + _contexto(ctx)
    marcas = {a: "[X]" if m["autoridad"] == a else "[ ]" for a in ("GG", "CA", "AG")}
    cuerpo += [_tabla(["Campo", "Detalle"], [
        ["Partida a modificar", f"{m['concepto']} — ciclo {m['ciclo']} del ejercicio {m['ejercicio']}"],
        ["Monto aprobado (Bs)", _bs(m["aprobado"])],
        ["Monto solicitado (Bs) y variación (%)", f"{_bs(m['solicitado'])} ({pct(m['variacion_pct'])})"],
        ["Motivo y desviación que la origina (FIN-TES-12)", m["motivo"]],
        ["Fuente de recursos", m["fuente_recursos"]],
        ["Efecto en el flujo de caja y en el SME", m["efecto_flujo"]],
        ["Autoridad que aprueba (según CI-FIN-09)",
         f"Gerencia General {marcas['GG']}   Consejo de Administración {marcas['CA']}   Asamblea General {marcas['AG']}"],
        ["Fundamento", "; ".join(m["fundamentos"])],
        ["Estado", f"{m['estado'].capitalize()}" + (f" el {fecha_pantalla(m['resuelto_en'])} por {m['resuelto_por']}"
                                                   if m.get("resuelto_en") else "")
         + (f" (acta {m['acta_ref']})" if m.get("acta_ref") else "")],
    ], [6.2, 11.2])]
    cuerpo += _firmas(["Solicitado por (Gerencia General)", "Verificado por (Unidad de Finanzas)",
                       "Autorizado por (según matriz CI-FIN-09)"],
                      [m["solicitado_por"], m.get("verificado_por") or "", m.get("resuelto_por") or ""])
    return _construir("FIN-TES-11", cuerpo, CICLO_PRESU)


def _titulo_ventana(c: Mapping[str, Any]) -> str:
    if c["tipo"] == "CICLO":
        return f"Ciclo {c['numero']} del ejercicio {c['ejercicio']}"
    if c["tipo"] == "SEMESTRE":
        return f"Semestre {c['numero']} del ejercicio {c['ejercicio']}"
    return f"Ejercicio {c['ejercicio']}"


def fin_tes_12(c: Mapping[str, Any], ctx: Contexto, acciones: Mapping[str, Mapping[str, Any]] | None = None) -> bytes:
    acciones = acciones or {}
    cuerpo = _encabezado("FIN-TES-12", f"Informe de Ejecución y Variaciones — {_titulo_ventana(c)}",
                         "Unidad de Finanzas", CICLO_PRESU) + _contexto(ctx)
    cuerpo += [_p(f"Período: {fecha_pantalla(c['desde'])} al {fecha_pantalla(c['hasta'])} · Verdes {c['verdes']} · "
                  f"amarillos {c['amarillos']} · rojos {c['rojos']}", E_SUB)]
    filas = []
    for f in c["filas"]:
        a = acciones.get(f["partida"]) or {}
        nota = NOMBRE_SEMAFORO[f["semaforo"]]
        if a:
            nota += f" — {a.get('causa', '')}" + (f" / {a['accion']}" if a.get("accion") else "")
        filas.append([f["nombre"], _bs(f["presupuesto"]), _bs(f["real"]), _bs(f["variacion"]), pct(f["pct"]), nota])
    if c.get("flujo"):
        fl = c["flujo"]
        a = acciones.get("FLUJO") or {}
        filas.append(["Saldo final de efectivo", _bs(fl["presupuesto"]), _bs(fl["real"]), _bs(fl["variacion"]),
                      pct(fl["pct"]), NOMBRE_SEMAFORO[fl["semaforo"]] + (f" — {a['causa']}" if a.get("causa") else "")])
    u = c.get("costo_unitario") or {}
    filas.append(["Costo unitario (Bs/kg)", num(u.get("presupuesto")), num(u.get("real")),
                  num(u.get("variacion")) if u.get("variacion") is not None else "—", pct(u.get("pct")),
                  NOMBRE_SEMAFORO.get(u.get("semaforo"))])
    cuerpo += [_tabla(["Partida", "Presupuesto (Bs)", "Real (Bs)", "Variación (Bs)", "Variación (%)",
                       "Semáforo y causa / acción"], filas, [4.3, 2.4, 2.4, 2.3, 1.8, 4.2], (1, 2, 3, 4))]
    precio = [f for f in c["filas"] if f.get("variacion_precio") is not None]
    if precio:
        cuerpo += [Spacer(1, 4), _p("Variaciones de precio y de cantidad", E_SUB),
                   _tabla(["Partida", "Cant. presup.", "Precio presup.", "Cant. real", "Precio real", "Var. precio",
                           "Var. cantidad"],
                          [[f["nombre"], num(f["cantidad_presupuesto"], 3), num(f["precio_presupuesto"], 4),
                            num(f["cantidad_real"], 3), num(f["precio_real"], 4), _bs(f["variacion_precio"]),
                            _bs(f["variacion_cantidad"])] for f in precio], [3.8, 2, 2, 2, 2, 2.8, 2.8],
                          (1, 2, 3, 4, 5, 6))]
    cuerpo += [Spacer(1, 4), _p("Toda variación en rojo exige causa, responsable y acción correctiva; una variación en "
                                "amarillo exige explicación escrita (Manual Presupuestario E.1).", E_PEQ)]
    cuerpo += _firmas(["Elaborado por (Unidad de Finanzas)", "Presentado por (Gerencia General)",
                       "Recibido por (Consejo de Administración)"])
    return _construir("FIN-TES-12", cuerpo, CICLO_PRESU)


def fin_tes_13(datos: Mapping[str, Any], ctx: Contexto, anterior: Mapping[str, Any] | None = None) -> bytes:
    cuerpo = _encabezado("FIN-TES-13", f"Cuadro de Mando de Indicadores — Ejercicio {datos['ejercicio']}",
                         "Unidad de Finanzas", CICLO_PRESU) + _contexto(ctx)
    previo = (anterior or {}).get("indicadores", {})
    filas = []
    for clave in ("ICL", "RINC", "ROA", "ICA", "MARGEN_OPERATIVO", "SME", "PMF", "DPF", "CAC", "MORTALIDAD"):
        i = datos["indicadores"].get(clave)
        if i is None:
            continue
        valor = num(i["valor"]) if clave == "ICA" else pct(i["valor"])
        prev = previo.get(clave, {}).get("valor")
        filas.append([i["nombre"], i["meta"], valor, NOMBRE_SEMAFORO.get(i["semaforo"]),
                      pct(i["proyectado"]) if i.get("proyectado") is not None else "—",
                      (num(prev) if clave == "ICA" else pct(prev)) if prev is not None else "—"])
    cuerpo += [_tabla(["Indicador", "Meta", "Resultado del período", "Semáforo", "Proyectado (E.3)",
                       "Período anterior"], filas, [4.6, 2.4, 2.6, 2, 2.6, 2.6], (2, 4, 5))]
    cuerpo += _firmas(["Elaborado por (Unidad de Finanzas)", "Verificado por (Consejo de Vigilancia, CI-FIN-18)",
                       "Recibido por (Gerencia General)"])
    return _construir("FIN-TES-13", cuerpo, CICLO_PRESU)


def inf_cv_02(informe: Mapping[str, Any], ctx: Contexto) -> bytes:
    d = informe["datos"]
    cuerpo = _encabezado("INF-CV-02", f"Informe Trimestral de Fiscalización de Liquidez — {informe['periodo']}",
                         "Consejo de Vigilancia", "Ciclo 1 — Tesorería (Financiamiento)") + _contexto(ctx)
    cuerpo += [_tabla(["Ciclo productivo", "SME vigente (Bs)", "Saldo real al cierre (Bs)", "% de cumplimiento",
                       "¿Cumple?"],
                      [[f"Ciclo {c['ciclo']}/{c['ejercicio']} ({fecha_pantalla(c['cierre'])})", _bs(c["sme"]),
                        _bs(c["saldo"]), pct(c["cumplimiento"]), "Sí" if c["cumple"] else "No"] for c in d["ciclos"]],
                      [5, 3, 3.4, 3, 3], (1, 2, 3)),
               Spacer(1, 4),
               _p(f"¿Se registran {d['minimo_reporte']} ciclos consecutivos de incumplimiento? "
                  f"{'Sí' if d['reportar_ag'] else 'No'}" + (f" — reportado a la Asamblea General el "
                                                             f"{fecha_pantalla(informe['reportado_ag_en'])}"
                                                             if informe.get("reportado_ag_en") else ""), E_SUB)]
    if informe.get("observaciones"):
        cuerpo += [_p(informe["observaciones"])]
    cuerpo += _firmas(["Elaborado por (Consejo de Vigilancia)", "Recibido por (Gerencia General)"],
                      [informe["elaborado_por"], ""])
    return _construir("INF-CV-02", cuerpo, "Ciclo 1 — Tesorería (Financiamiento)")


def inf_cv_07(informe: Mapping[str, Any], ctx: Contexto) -> bytes:
    d = informe["datos"]
    r = d.get("resultados", {})
    marca = lambda ok: "Conforme [X]  Con observaciones [ ]" if ok else "Conforme [ ]  Con observaciones [X]"  # noqa: E731
    cuerpo = _encabezado("INF-CV-07", f"Informe de Fiscalización Mensual — {d.get('periodo', informe['periodo'])}",
                         "Consejo de Vigilancia", CICLO_INF) + _contexto(ctx)
    cuerpo += [_tabla(["Verificación (CI-FIN-18)", "Resultado"], [
        ["Partidas muestreadas del Consolidado de Costeo (FIN-INF-03) con respaldo razonable", marca(r.get("partidas"))],
        ["Recálculo independiente de ICL, RINC y ROA coincide con lo reportado", marca(r.get("recalculo"))],
        ["Razonabilidad del margen bruto frente al mes anterior y al presupuesto", marca(r.get("razonabilidad"))],
    ], [11, 6.4]), Spacer(1, 4), _p("Partidas muestreadas", E_SUB)]
    cuerpo += [_p(f"• {m}") for m in d.get("muestra", [])]
    cuerpo += [Spacer(1, 4), _p("Recálculo independiente (acumulado del ejercicio)", E_SUB),
               _tabla(["Indicador", "Reportado (FIN-INF-04)", "Recalculado por el CV"],
                      [[k, pct(d["reportado"].get(k)), pct(d["recalculo"].get(k))] for k in ("ICL", "RINC", "ROA")],
                      [5, 6.2, 6.2], (1, 2)),
               _p(f"Margen bruto del mes {pct(d.get('margen_bruto_mes'))}; mes anterior "
                  f"{pct(d.get('margen_bruto_mes_anterior'))}; presupuesto {pct(d.get('margen_bruto_presupuesto'))}.",
                  E_PEQ),
               _p("¿Se emite ANTES de la comunicación oficial de indicadores? Sí [X]   No [ ] (CI-FIN-18)", E_SUB)]
    if informe.get("observaciones"):
        cuerpo += [_p("Observaciones", E_SUB), _p(informe["observaciones"])]
    cuerpo += _firmas(["Elaborado por (Consejo de Vigilancia)", "Recibido por (Gerencia General, antes del FIN-INF-05)"],
                      [informe["elaborado_por"], ""])
    return _construir("INF-CV-07", cuerpo, CICLO_INF)


def fin_inf_05(informe: Mapping[str, Any], ctx: Contexto) -> bytes:
    cuerpo = _encabezado("FIN-INF-05", f"Informe Gerencial de Indicadores — {informe['periodo']}", "Gerencia General",
                         CICLO_INF) + _contexto(ctx)
    cuerpo += [_tabla(["Indicador", "Resultado", "Umbral alcanzado", "Desviación significativa",
                       "Acción / decisión gerencial"],
                      [[k, pct(v["valor"]), (v["umbral"] or "—").capitalize(), "Sí" if v["desviacion"] else "No",
                        v.get("accion") or "—"] for k, v in informe["indicadores"].items()],
                      [2.4, 2.4, 3, 3, 6.6], (1,))]
    if informe.get("observaciones_incorporadas"):
        cuerpo += [Spacer(1, 4), _p("Observaciones del INF-CV-07 incorporadas", E_SUB),
                   _p(informe["observaciones_incorporadas"])]
    cuerpo += [Spacer(1, 4), _p("Emitido una vez recibido el INF-CV-07 del mismo período (CI-FIN-18, Paso 6).", E_PEQ)]
    cuerpo += _firmas(["Elaborado por (Gerencia General)", "Con observaciones incorporadas de (Consejo de Vigilancia, "
                                                           "INF-CV-07)"], [informe["emitido_por"], ""])
    return _construir("FIN-INF-05", cuerpo, CICLO_INF)


def _hoja(ws: Any, titulo: str, encabezados: Sequence[str], filas: Sequence[Sequence[Any]],
          anchos: Sequence[int]) -> None:
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(list(encabezados))
    relleno = PatternFill("solid", fgColor="E3EEE7")
    for celda in ws[2]:
        celda.font = Font(bold=True)
        celda.fill = relleno
        celda.alignment = Alignment(wrap_text=True)
    for fila in filas:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for fila in ws.iter_rows(min_row=3):
        for celda in fila:
            if isinstance(celda.value, float):
                celda.number_format = "#,##0.00"


def excel_maestro(series: Mapping[str, Sequence[Mapping[str, Any]]], nombres: Mapping[str, str]) -> bytes:
    """Presupuesto Maestro 2026-2030: una hoja por módulo."""
    wb = Workbook()
    wb.remove(wb.active)
    anios = ["2026", "2027", "2028", "2029", "2030"]
    for modulo, filas in series.items():
        ws = wb.create_sheet(modulo[:31])
        _hoja(ws, f"Presupuesto Maestro 2026-2030 — {nombres.get(modulo, modulo)}", ["Concepto", "Unidad", *anios],
              [[f["concepto"], f.get("unidad") or "", *[f.get(a) for a in anios]] for f in filas], [52, 10, 15, 15, 15,
                                                                                                   15, 15])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def excel_fin_tes_12(c: Mapping[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FIN-TES-12"
    filas = [[f["nombre"], f["presupuesto"], f["real"], f["variacion"], float(f["pct"]) if f["pct"] is not None else None,
              f["semaforo"], f.get("variacion_precio"), f.get("variacion_cantidad")] for f in c["filas"]]
    if c.get("flujo"):
        fl = c["flujo"]
        filas.append(["Saldo final de efectivo", fl["presupuesto"], fl["real"], fl["variacion"],
                      float(fl["pct"]) if fl["pct"] is not None else None, fl["semaforo"], None, None])
    _hoja(ws, f"FIN-TES-12 — {_titulo_ventana(c)}", ["Partida", "Presupuesto (Bs)", "Real (Bs)", "Variación (Bs)",
                                                    "Variación (fracción)", "Semáforo", "Var. precio", "Var. cantidad"],
          filas, [40, 16, 16, 16, 14, 12, 14, 14])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
