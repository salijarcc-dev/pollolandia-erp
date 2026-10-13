"""Pruebas de la Fase 12: formatos PDF y Excel (reportes/formatos_presupuesto.py).

FIN-TES-09 a FIN-TES-13, INF-CV-02, INF-CV-07 y FIN-INF-05 se generan con las
cifras del escenario 2026 y llevan el pie «Página 1 de N»; el Excel del
Presupuesto Maestro tiene una hoja por módulo.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from pruebas import escenario_2026 as esc
from reportes import formatos_presupuesto as rep
from reportes.formatos_tesoreria import Contexto
from servicios import cierre, indicadores, presupuesto, vigilancia

D = Decimal
CTX = Contexto(fecha=date(2026, 10, 12), elaborado_por="Jefe de Finanzas", cargo="Jefe de Finanzas", ciclo="")
MUESTRA = ["Costo de ventas de carne 6.1.02.01", "6.1.01.01 Centro de costos carne", "5.1.01.01 Venta de carne"]


def _es_pdf(datos: bytes) -> None:
    assert datos.startswith(b"%PDF") and b"P\\341gina 1 de" in datos


@pytest.fixture()
def cargado(con, s):
    esc.cargar_escenario_2026(con, s)
    return con


def test_formatos_numericos():
    assert rep.pct(D("0.6241")) == "62,4 %" and rep.pct(None) == "—"
    assert rep.num(D("1234.5")) == "1.234,50" and rep.num(D("2.0"), 0) == "2"


def test_pdf_del_presupuesto(cargado, s):
    con = cargado
    _es_pdf(rep.fin_tes_09(2026, presupuesto.hoja_supuestos(con, s["jf"], 2026), CTX))
    _es_pdf(rep.fin_tes_10(presupuesto.presupuesto_ciclo(con, s["jf"], 2026, 1), CTX))
    calc = presupuesto.calcular_ejecucion(con, s["jf"], 2026, "CICLO", 1)
    _es_pdf(rep.fin_tes_12(calc, CTX))
    m = presupuesto.solicitar_modificacion(con, s["gg"], 2026, 1, "GASTOS_VENTA", "3100", motivo="Flete adicional",
                                           fuente_recursos="Recursos propios", efecto_flujo="Sin efecto en el SME",
                                           fecha="2026-08-01")
    mod = next(x for x in presupuesto.listar_modificaciones(con, s["jf"], 2026) if x["id"] == m["id"])
    pdf = rep.fin_tes_11(mod, CTX)
    _es_pdf(pdf)
    ind = indicadores.indicadores_ejercicio(con, s["jf"], 2026)
    _es_pdf(rep.fin_tes_13(ind, CTX))


def test_excel_maestro_y_fin_tes_12(cargado, s):
    con = cargado
    series = {m: presupuesto.serie_maestro(con, s["jf"], m) for m in presupuesto.MODULOS_MAESTRO}
    wb = load_workbook(io.BytesIO(rep.excel_maestro(series, presupuesto.MODULOS_MAESTRO)))
    assert len(wb.sheetnames) == len(presupuesto.MODULOS_MAESTRO)
    calc = presupuesto.calcular_ejecucion(con, s["jf"], 2026, "CICLO", 1)
    hoja = load_workbook(io.BytesIO(rep.excel_fin_tes_12(calc))).active
    valores = [c.value for fila in hoja.iter_rows() for c in fila]
    assert 183300.0 in valores                                        # ingresos por carne reales del ciclo 1


def test_pdf_de_vigilancia_e_informe_gerencial(cargado, s):
    con = cargado
    sep = int(con.execute("SELECT id FROM periodos WHERE ejercicio = 2026 AND numero = 3").fetchone()[0])
    vigilancia.registrar_inf_cv_02(con, s["cv"], 2026, 1, "2026-10-03")
    _es_pdf(rep.inf_cv_02(vigilancia.listar_informes_cv(con, s["cv"], "INF-CV-02")[0], CTX))
    cierre.registrar_lista_verificacion(con, s["jf"], sep, "2026-10-02")
    indicadores.registrar_inf_cv_07(con, s["cv"], sep, "2026-10-05", partidas_muestreadas=MUESTRA,
                                    partidas_conformes=True, recalculo_conforme=True, razonabilidad_conforme=True)
    _es_pdf(rep.inf_cv_07(vigilancia.listar_informes_cv(con, s["cv"], "INF-CV-07")[0], CTX))
    indicadores.emitir_fin_inf_05(con, s["gg"], sep, "2026-10-08", decisiones={
        "RINC": {"accion": "Previsto en el presupuesto"}, "ROA": {"accion": "Se evalúa al cierre anual"}})
    _es_pdf(rep.fin_inf_05(indicadores.listar_fin_inf_05(con, s["gg"])[0], CTX))
